from pathlib import Path
from shutil import copy2
import csv

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Image, PageBreak, Spacer
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics


ROOT = Path(__file__).resolve().parents[1]
RUN_DIR = ROOT / "runs" / "detect" / "runs" / "detect" / "runs" / "fasdd_uav_yolov8n_v32_print_demo_hardneg"
DOC_DIR = ROOT / "文档" / "模型评估" / "Word"
PDF_DIR = ROOT / "文档" / "模型评估" / "PDF"
ASSET_DIR = DOC_DIR / "评估图表_v4"

for directory in (DOC_DIR, PDF_DIR, ASSET_DIR):
    directory.mkdir(parents=True, exist_ok=True)

FIGURES = [
    ("labels.jpg", "图 1  标签数量与目标框分布"),
    ("confusion_matrix.png", "图 2  原始混淆矩阵"),
    ("confusion_matrix_normalized.png", "图 3  归一化混淆矩阵"),
    ("results.png", "图 4  训练指标与损失曲线"),
    ("BoxPR_curve.png", "图 5  Precision-Recall 曲线"),
    ("BoxF1_curve.png", "图 6  F1-Confidence 曲线"),
    ("BoxP_curve.png", "图 7  Precision-Confidence 曲线"),
    ("BoxR_curve.png", "图 8  Recall-Confidence 曲线"),
]


def copy_figures():
    paths = []
    for filename, caption in FIGURES:
        src = RUN_DIR / filename
        if src.exists():
            dst = ASSET_DIR / filename
            copy2(src, dst)
            paths.append((dst, caption))
    return paths


with (RUN_DIR / "results.csv").open("r", encoding="utf-8") as f:
    LAST = list(csv.DictReader(f))[-1]


def val(key, digits=5):
    return f"{float(LAST[key]):.{digits}f}"


def pct(key):
    return f"{float(LAST[key]) * 100:.2f}%"


CORE_METRICS = [
    ["Precision", val("metrics/precision(B)"), pct("metrics/precision(B)"), "误报控制表现较好，模型输出的检测结果具有较高可信度。"],
    ["Recall", val("metrics/recall(B)"), pct("metrics/recall(B)"), "目标检出能力较稳定，能够覆盖大多数真实火焰与烟雾目标。"],
    ["mAP50", val("metrics/mAP50(B)"), pct("metrics/mAP50(B)"), "综合检测精度较高，满足常规目标检测评价要求。"],
    ["mAP50-95", val("metrics/mAP50-95(B)"), pct("metrics/mAP50-95(B)"), "严格定位评价下仍有提升空间，后续可继续优化检测框质量。"],
]

LOSS_ROWS = [
    ["train/box_loss", val("train/box_loss"), "训练集检测框回归损失"],
    ["train/cls_loss", val("train/cls_loss"), "训练集分类损失"],
    ["train/dfl_loss", val("train/dfl_loss"), "训练集分布焦点损失"],
    ["val/box_loss", val("val/box_loss"), "验证集检测框回归损失"],
    ["val/cls_loss", val("val/cls_loss"), "验证集分类损失"],
    ["val/dfl_loss", val("val/dfl_loss"), "验证集分布焦点损失"],
]

CHECK_ROWS = [
    ["历史误报样本池回归测试", "322 张", "2 张误报", "0.62%"],
    ["综合负样本回归测试", "338 张", "2 张误报", "0.59%"],
    ["验证集目标级正确识别", "16007 个", "fire 10565 / smoke 5442", "来自混淆矩阵"],
]

SUMMARY = (
    "v4.0_for_preview 模型在最终验证结果中取得 Precision 0.88217、Recall 0.86460、"
    "mAP50 0.91846、mAP50-95 0.66894。结果表明，模型在常规 IoU 条件下具备较高检测精度，"
    "误报控制和目标检出能力较为均衡；在更严格的多阈值定位评价下仍存在进一步提升空间。"
)


def set_docx_font(run, size=None, bold=False):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    run.bold = bold


def add_docx_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text))
        set_docx_font(r, 9.5, True)

    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 3 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(text))
            set_docx_font(r, 9)

    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def add_docx_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_docx_font(r, 10.5)
    return p


def build_docx(figure_paths):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    for style_name in ["Normal", "Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("林焰智巡 v4.0_for_preview 模型性能评估报告")
    set_docx_font(r, 20, True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Fire and Smoke Detection Model Performance Report")
    set_docx_font(r, 10)
    r.font.color.rgb = RGBColor(80, 80, 80)

    doc.add_heading("一、报告摘要", level=1)
    add_docx_paragraph(doc, SUMMARY)

    add_docx_table(
        doc,
        ["项目", "内容"],
        [
            ["模型版本", "v4.0_for_preview"],
            ["任务类型", "火焰 / 烟雾目标检测"],
            ["类别定义", "fire、smoke"],
            ["输入尺度", "640"],
            ["默认置信度阈值", "0.25"],
        ],
        [1.6, 4.8],
    )

    doc.add_heading("二、核心性能指标", level=1)
    add_docx_table(doc, ["指标", "数值", "百分比", "评估"], CORE_METRICS, [1.2, 1.0, 1.0, 4.2])
    add_docx_paragraph(
        doc,
        "mAP50 达到 91.85%，说明模型在常规目标检测评价条件下具有较高识别精度。"
        "Precision 为 88.22%，Recall 为 86.46%，说明模型在减少误报和保持目标检出之间取得了较均衡的表现。"
        "mAP50-95 为 66.89%，表明在更严格的边界框定位评价下，模型仍有进一步优化空间。",
    )

    doc.add_heading("三、补充测试结果", level=1)
    add_docx_table(doc, ["测试项", "样本规模", "结果", "评估"], CHECK_ROWS, [2.3, 1.2, 1.8, 1.3])
    add_docx_paragraph(
        doc,
        "补充测试用于观察模型对误报样本的回归表现。历史误报样本池误报率为 0.62%，综合负样本回归测试误报率为 0.59%。"
        "该结果说明模型对已收集干扰样本具有较好的误报抑制效果。",
    )

    doc.add_heading("四、损失收敛情况", level=1)
    add_docx_table(doc, ["损失项", "最终值", "说明"], LOSS_ROWS, [1.8, 1.3, 3.3])
    add_docx_paragraph(
        doc,
        "训练损失和验证损失用于判断模型训练过程是否稳定。box_loss 反映检测框回归误差，"
        "cls_loss 反映类别分类误差，dfl_loss 反映边界框分布回归误差。",
    )

    doc.add_heading("五、混淆矩阵与误差分析", level=1)
    add_docx_paragraph(
        doc,
        "原始混淆矩阵显示，fire 与 smoke 之间的相互混淆数量较低，模型对两类目标具有较好的区分能力。"
        "主要误差集中在目标与 background 的边界判断上，即部分真实目标未被检出，或部分背景区域被识别为目标。"
        "归一化混淆矩阵进一步展示了各类别在比例尺度下的识别稳定性。",
    )

    doc.add_heading("六、标签分布与样本结构", level=1)
    add_docx_paragraph(
        doc,
        "标签分布图显示，fire 类样本数量高于 smoke 类样本，类别分布存在一定不均衡。"
        "目标框分布主要集中在画面中部区域，小尺寸目标占比较高，这对模型的小目标检测能力提出了要求。",
    )

    doc.add_heading("七、性能曲线分析", level=1)
    add_docx_paragraph(
        doc,
        "Precision-Recall 曲线用于观察精确率与召回率之间的权衡关系；F1-Confidence 曲线用于确定综合性能较优的置信度区间；"
        "Precision-Confidence 与 Recall-Confidence 曲线分别反映置信度阈值变化对误报控制和目标检出的影响。",
    )

    doc.add_heading("八、推理性能", level=1)
    add_docx_table(
        doc,
        ["项目", "结果"],
        [
            ["FP RKNN 样例级应用推理", "约 42 ms/帧"],
            ["计时输入", "1280x720 图像，经 letterbox 进入 640 尺度模型"],
            ["稳定速度参考", "FP RKNN"],
            ["INT8 量化状态", "已生成，但不作为当前稳定性能结论"],
        ],
        [2.4, 4.0],
    )
    add_docx_paragraph(doc, "推理速度用于衡量模型端侧运行效率。实际运行速度会受到输入分辨率、预处理、后处理和硬件负载影响。")

    doc.add_heading("九、模型性能图", level=1)
    for path, caption in figure_paths:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        width = 5.7 if path.name in ["labels.jpg", "confusion_matrix.png", "confusion_matrix_normalized.png"] else 6.1
        p.add_run().add_picture(str(path), width=Inches(width))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        set_docx_font(r, 9)

    doc.add_heading("十、结论", level=1)
    add_docx_paragraph(
        doc,
        SUMMARY
        + " 后续优化重点应放在提升 mAP50-95、改善严格定位条件下的检测框质量，并继续扩大负样本覆盖范围以进一步降低误报。",
    )

    output = DOC_DIR / "v4_model_performance_report_fixed.docx"
    doc.save(output)
    return output


def build_pdf(figure_paths):
    pdfmetrics.registerFont(TTFont("SimHeiLocal", r"C:\Windows\Fonts\simhei.ttf"))
    base = "SimHeiLocal"
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("CNTitle", fontName=base, fontSize=20, leading=26, alignment=1, spaceAfter=8))
    styles.add(ParagraphStyle("CNSub", fontName=base, fontSize=9, leading=12, alignment=1, textColor=colors.HexColor("#555555"), spaceAfter=10))
    styles.add(ParagraphStyle("CNH1", fontName=base, fontSize=13, leading=18, spaceBefore=8, spaceAfter=6))
    styles.add(ParagraphStyle("CNBody", fontName=base, fontSize=9.2, leading=14, firstLineIndent=18, spaceAfter=5))
    styles.add(ParagraphStyle("CNCaption", fontName=base, fontSize=8.2, leading=10, alignment=1, spaceAfter=8))

    def paragraph(text):
        return Paragraph(str(text), styles["CNBody"])

    def table(headers, rows, widths):
        data = [[Paragraph(str(x), styles["CNBody"]) for x in headers]]
        data += [[Paragraph(str(x), styles["CNBody"]) for x in row] for row in rows]
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), base),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.black),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        return t

    story = [
        Paragraph("林焰智巡 v4.0_for_preview 模型性能评估报告", styles["CNTitle"]),
        Paragraph("Fire and Smoke Detection Model Performance Report", styles["CNSub"]),
        Paragraph("一、报告摘要", styles["CNH1"]),
        paragraph(SUMMARY),
        table(
            ["项目", "内容"],
            [
                ["模型版本", "v4.0_for_preview"],
                ["任务类型", "火焰 / 烟雾目标检测"],
                ["类别定义", "fire、smoke"],
                ["输入尺度", "640"],
                ["默认置信度阈值", "0.25"],
            ],
            [38 * mm, 115 * mm],
        ),
        Paragraph("二、核心性能指标", styles["CNH1"]),
        table(["指标", "数值", "百分比", "评估"], CORE_METRICS, [32 * mm, 25 * mm, 25 * mm, 75 * mm]),
        paragraph("mAP50 达到 91.85%，说明模型在常规目标检测评价条件下具有较高识别精度。Precision 为 88.22%，Recall 为 86.46%，说明模型在减少误报和保持目标检出之间取得了较均衡的表现。mAP50-95 为 66.89%，表明在更严格的边界框定位评价下，模型仍有进一步优化空间。"),
        Paragraph("三、补充测试结果", styles["CNH1"]),
        table(["测试项", "样本规模", "结果", "评估"], CHECK_ROWS, [54 * mm, 28 * mm, 42 * mm, 32 * mm]),
        paragraph("补充测试用于观察模型对误报样本的回归表现。历史误报样本池误报率为 0.62%，综合负样本回归测试误报率为 0.59%。该结果说明模型对已收集干扰样本具有较好的误报抑制效果。"),
        Paragraph("四、损失收敛情况", styles["CNH1"]),
        table(["损失项", "最终值", "说明"], LOSS_ROWS, [45 * mm, 30 * mm, 80 * mm]),
        paragraph("训练损失和验证损失用于判断模型训练过程是否稳定。box_loss 反映检测框回归误差，cls_loss 反映类别分类误差，dfl_loss 反映边界框分布回归误差。"),
        PageBreak(),
        Paragraph("五、混淆矩阵与误差分析", styles["CNH1"]),
        paragraph("原始混淆矩阵显示，fire 与 smoke 之间的相互混淆数量较低，模型对两类目标具有较好的区分能力。主要误差集中在目标与 background 的边界判断上，即部分真实目标未被检出，或部分背景区域被识别为目标。归一化混淆矩阵进一步展示了各类别在比例尺度下的识别稳定性。"),
        Paragraph("六、标签分布与样本结构", styles["CNH1"]),
        paragraph("标签分布图显示，fire 类样本数量高于 smoke 类样本，类别分布存在一定不均衡。目标框分布主要集中在画面中部区域，小尺寸目标占比较高，这对模型的小目标检测能力提出了要求。"),
        Paragraph("七、性能曲线分析", styles["CNH1"]),
        paragraph("Precision-Recall 曲线用于观察精确率与召回率之间的权衡关系；F1-Confidence 曲线用于确定综合性能较优的置信度区间；Precision-Confidence 与 Recall-Confidence 曲线分别反映置信度阈值变化对误报控制和目标检出的影响。"),
        Paragraph("八、推理性能", styles["CNH1"]),
        table(
            ["项目", "结果"],
            [
                ["FP RKNN 样例级应用推理", "约 42 ms/帧"],
                ["计时输入", "1280x720 图像，经 letterbox 进入 640 尺度模型"],
                ["稳定速度参考", "FP RKNN"],
                ["INT8 量化状态", "已生成，但不作为当前稳定性能结论"],
            ],
            [58 * mm, 95 * mm],
        ),
        paragraph("推理速度用于衡量模型端侧运行效率。实际运行速度会受到输入分辨率、预处理、后处理和硬件负载影响。"),
        PageBreak(),
        Paragraph("九、模型性能图", styles["CNH1"]),
    ]

    for path, caption in figure_paths:
        width = 150 * mm if path.name in ["labels.jpg", "confusion_matrix.png", "confusion_matrix_normalized.png"] else 165 * mm
        ratio = 0.72 if path.name == "labels.jpg" else (0.78 if "confusion" in path.name else 0.62)
        img = Image(str(path), width=width, height=width * ratio)
        img.hAlign = "CENTER"
        story.extend([img, Paragraph(caption, styles["CNCaption"])])

    story.extend(
        [
            Paragraph("十、结论", styles["CNH1"]),
            paragraph(SUMMARY + " 后续优化重点应放在提升 mAP50-95、改善严格定位条件下的检测框质量，并继续扩大负样本覆盖范围以进一步降低误报。"),
        ]
    )

    output = PDF_DIR / "v4_model_performance_report_fixed.pdf"
    SimpleDocTemplate(
        str(output),
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    ).build(story)
    return output


if __name__ == "__main__":
    figures = copy_figures()
    print(build_docx(figures))
    print(build_pdf(figures))
    print(f"figures={len(figures)}")


