from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
src = ROOT / "文档" / "模型评估" / "Markdown" / "model_training_defense_guide.md"
out = ROOT / "文档" / "模型评估" / "Word" / "model_training_defense_guide.docx"
out.parent.mkdir(parents=True, exist_ok=True)

text = src.read_text(encoding="utf-8")

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

for style_name in ["Normal", "Heading 1", "Heading 2"]:
    style = doc.styles[style_name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Heading 1"].font.size = Pt(15)
doc.styles["Heading 2"].font.size = Pt(12)


def set_font(run, size=None, bold=False):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    run.bold = bold


def add_para(value, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    r = p.add_run(value)
    set_font(r, 10.5)


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_font(r, 9.5, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_font(r, 9)
    doc.add_paragraph()


lines = text.splitlines()
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line:
        i += 1
        continue
    if line.startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line[2:].strip())
        set_font(r, 20, True)
        i += 1
        continue
    if line.startswith("## "):
        doc.add_heading(line[3:].strip(), level=1)
        i += 1
        continue
    if line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=2)
        i += 1
        continue
    if line.startswith("| "):
        table_lines = []
        while i < len(lines) and lines[i].startswith("| "):
            table_lines.append(lines[i])
            i += 1
        parsed = [[cell.strip() for cell in row.strip("|").split("|")] for row in table_lines]
        if len(parsed) >= 2:
            headers = parsed[0]
            rows = [row for row in parsed[2:] if len(row) == len(headers)]
            add_table(headers, rows)
        continue
    if line[0:3].replace(".", "").isdigit() and ". " in line[:4]:
        add_para(line, indent=False)
        i += 1
        continue
    if line.startswith("- "):
        add_para(line[2:], indent=False)
        i += 1
        continue
    add_para(line)
    i += 1

doc.save(out)
print(out)
