from pathlib import Path
from shutil import copy2

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


path = Path("文档/模型评估/Word/v4_model_performance_report.docx")
backup = Path("文档/模型评估/Word/v4_model_performance_report_before_image_descriptions.docx")

if path.exists() and not backup.exists():
    copy2(path, backup)

doc = Document(path)

descriptions = {
    "图 1  标签数量与目标框分布": "图示说明：该图展示训练数据中 fire 与 smoke 两类标签的数量分布，以及目标框在画面中的位置和尺寸分布。fire 类样本数量高于 smoke 类样本，目标框主要集中在画面中部，且小尺寸目标占比较高，说明模型训练需要兼顾类别不均衡和小目标检测能力。",
    "图 2  原始混淆矩阵": "图示说明：原始混淆矩阵反映各类别在验证集上的具体识别数量。fire 与 smoke 的相互混淆数量较低，说明模型对两类目标具备较好的区分能力；主要误差来自目标与 background 的边界判断。",
    "图 3  归一化混淆矩阵": "图示说明：归一化混淆矩阵以比例形式展示各类别识别结果，便于比较不同类别在样本数量不一致时的识别稳定性。该图可用于观察模型在 fire、smoke 与 background 之间的误判比例。",
    "图 4  训练指标与损失曲线": "图示说明：该图展示训练过程中损失和验证指标的变化趋势。最终指标稳定在 Precision 0.882、Recall 0.865、mAP50 0.918、mAP50-95 0.669 附近，说明模型训练过程收敛较稳定。",
    "图 5  Precision-Recall 曲线": "图示说明：PR 曲线用于评估模型在不同阈值下精确率与召回率的权衡关系。曲线整体面积越大，表示模型在保持检出能力的同时能够更好地控制误报。",
    "图 6  F1-Confidence 曲线": "图示说明：F1-Confidence 曲线用于观察不同置信度阈值下 Precision 与 Recall 的综合表现，可作为选择默认置信度阈值的参考。",
    "图 7  Precision-Confidence 曲线": "图示说明：该曲线反映置信度阈值升高时精确率的变化趋势。通常阈值越高，误报越少，但过高阈值可能降低召回能力。",
    "图 8  Recall-Confidence 曲线": "图示说明：该曲线反映置信度阈值升高时召回率的变化趋势。阈值过高会导致部分真实目标被过滤，因此需要结合 Precision 曲线共同选择阈值。",
}

inserted = 0
for para in list(doc.paragraphs):
    caption = para.text.strip()
    if caption not in descriptions:
        continue

    next_el = para._p.getnext()
    if next_el is not None:
        next_text = "".join(next_el.itertext()).strip()
        if next_text.startswith("图示说明："):
            continue

    new_para = doc.add_paragraph()
    para._p.addnext(new_para._p)
    run = new_para.add_run(descriptions[caption])
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(60, 60, 60)
    new_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    new_para.paragraph_format.first_line_indent = Pt(19)
    new_para.paragraph_format.space_after = Pt(6)
    inserted += 1

try:
    doc.save(path)
    print(path.resolve())
except PermissionError:
    fallback = Path("output/doc/v4_model_performance_report_with_image_descriptions.docx")
    doc.save(fallback)
    print(fallback.resolve())

print(f"inserted={inserted}")
if backup.exists():
    print(f"backup={backup.resolve()}")
